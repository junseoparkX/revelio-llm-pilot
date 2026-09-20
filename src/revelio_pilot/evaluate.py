from __future__ import annotations

import argparse
import ast
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .schema import validate_prediction


LABEL_ORDER = ["YES", "NO", "UNCERTAIN"]
GROUP_ORDER = ["SEO_ONLY", "GEO_ONLY", "BOTH", "NEITHER", "UNCERTAIN"]

FIELD_DEFINITIONS = [
    ("seo_duty", "Does the job assign current SEO work?", "YES, NO, or UNCERTAIN"),
    ("geo_duty", "Does the job assign current AI-search visibility work (GEO/AEO)?", "YES, NO, or UNCERTAIN"),
    ("seo_centrality", "How important is SEO in the role?", "PRIMARY, SECONDARY, NOT_APPLICABLE, or UNCLEAR"),
    ("geo_centrality", "How important is GEO/AEO in the role?", "PRIMARY, SECONDARY, NOT_APPLICABLE, or UNCLEAR"),
    ("seo_evidence", "Exact words from the posting that support the SEO decision.", "Blank when no evidence is present"),
    ("geo_evidence", "Exact words from the posting that support the GEO/AEO decision.", "Blank when no evidence is present"),
    ("required_prior_experience", "Prior experience requested by the employer, kept separate from current duties.", "Short text summary"),
    ("prior_experience_evidence", "Exact words supporting the prior-experience summary.", "Blank when none is stated"),
    ("seo_background_for_geo", "Whether the posting connects SEO background to current GEO work.", "EXPLICIT, SUGGESTIVE, or NO_EVIDENCE"),
    ("adjacent_type", "Nearby but excluded work, such as paid search or internal product search.", "Short category or blank"),
    ("text_completeness", "Whether enough posting text was available to judge the role.", "FULL, PARTIAL, or UNREADABLE"),
    ("uncertainty_reason", "Why a duty could not be decided.", "Required when either duty is UNCERTAIN"),
    ("summary_group", "One combined label for the posting.", "SEO_ONLY, GEO_ONLY, BOTH, NEITHER, or UNCERTAIN"),
    ("concise_rationale", "A short explanation of the decision.", "One to three sentences"),
]

EXCEL_COLUMN_LABELS = {
    "model_name": "Model", "n_scored": "Postings scored", "job_id": "Job ID",
    "title_raw": "Original title", "jobtitle_translated": "Translated title",
    "reference_origin": "Reference source", "pilot_stratum": "Sample group",
    "seo_true": "Reference SEO", "seo_pred": "Model SEO",
    "geo_true": "Reference AI search", "geo_pred": "Model AI search",
    "group_true": "Reference combined group", "group_pred": "Model combined group",
    "seo_evidence": "SEO evidence quote", "geo_evidence": "AI-search evidence quote",
    "seo_centrality": "SEO importance", "geo_centrality": "AI-search importance",
    "required_prior_experience": "Requested prior experience",
    "prior_experience_evidence": "Prior-experience evidence quote",
    "seo_background_for_geo": "SEO background linked to AI-search work",
    "adjacent_type": "Nearby excluded work", "text_completeness": "Posting text quality",
    "uncertainty_reason": "Reason for uncertainty", "concise_rationale": "Model explanation",
    "seo_evidence_valid": "SEO quotation found", "geo_evidence_valid": "AI-search quotation found",
    "seo_precision": "SEO precision", "seo_recall": "SEO recall", "seo_f1": "SEO F1",
    "geo_precision": "AI-search precision", "geo_recall": "AI-search recall", "geo_f1": "AI-search F1",
    "group_accuracy": "Combined-group accuracy", "estimated_cost_usd": "Estimated cost (USD)",
    "mean_latency_seconds": "Mean time (seconds)",
    "possible_duty_prior_confusions": "Possible duty/experience mix-ups",
    "possible_duty_prior_confusion": "Possible duty/experience mix-up",
    "input_tokens": "Input tokens", "output_tokens": "Output tokens",
    "latency_seconds": "Time (seconds)", "attempts": "Attempts",
    "successful_calls": "Successful calls", "terminal_failed_calls": "Failed calls",
    "terminal_invalid_output_calls": "Invalid outputs", "failed_call_rate": "Failed-call rate",
    "invalid_output_rate": "Invalid-output rate", "schema_or_logic_valid": "Reference structure valid",
    "evidence_substrings_valid": "Reference quotations valid",
    "schema_or_logic_errors": "Reference structure problems", "evidence_errors": "Reference quotation problems",
    "plain_language_meaning": "Meaning", "values_or_format": "Allowed values or format", "field": "Field name",
    "true_label": "Reference answer", "predicted_label": "Model answer", "count": "Number of postings",
}


def _clean_excel_value(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, bool) or type(value).__name__ == "bool_":
        return "Yes" if bool(value) else "No"
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _write_frame_sheet(workbook: Workbook, name: str, frame: pd.DataFrame) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    frame = frame.rename(columns=EXCEL_COLUMN_LABELS)
    columns = list(frame.columns)
    if not columns:
        sheet["A1"] = "No records"
        return
    sheet.append(columns)
    for row in frame.itertuples(index=False, name=None):
        sheet.append([_clean_excel_value(value) for value in row])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    sheet.freeze_panes = "G2" if name in {"Posting results", "Disagreements"} else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if sheet.max_row > 1:
        safe_name = "".join(character for character in name.title() if character.isalnum()) + "Table"
        table = Table(displayName=safe_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    for index, column in enumerate(columns, start=1):
        values = [str(column)] + [str(_clean_excel_value(value)) for value in frame.iloc[:, index - 1].head(200)]
        width = min(max(max(map(len, values)) + 2, 11), 42)
        sheet.column_dimensions[get_column_letter(index)].width = width
        if column.endswith(" rate") or column in {
            "SEO precision", "SEO recall", "SEO F1", "AI-search precision", "AI-search recall",
            "AI-search F1", "Combined-group accuracy",
        }:
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = "0.0%"
        elif column == "Estimated cost (USD)":
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = '"$"0.0000'
        elif column in {"Time (seconds)", "Mean time (seconds)"}:
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = "0.00"
    sheet.row_dimensions[1].height = 30
    if name == "Field guide":
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 72
        sheet.column_dimensions["C"].width = 54
        for row in sheet.iter_rows(min_row=2, max_col=3):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 32


def write_excel_workbook(
    output: Path,
    per_model: pd.DataFrame,
    scored: pd.DataFrame,
    disagreements: pd.DataFrame,
    call_quality: pd.DataFrame,
    confusion: pd.DataFrame,
    reference_quality: pd.DataFrame,
) -> None:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.sheet_view.showGridLines = False
    overview["A2"] = "LLM pilot results"
    overview["A2"].font = Font(name="Arial", size=15, bold=True, color="1F1F1F")
    overview["A3"] = "This workbook compares how the candidate models classified current SEO and AI-search duties in job postings."
    overview["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
    overview["A5"] = "How to read this file"
    overview["A5"].font = Font(name="Arial", size=11, bold=True, color="1F4E78")
    guidance = [
        "Model summary shows overall accuracy, missed positive duties, speed, and estimated API cost.",
        "Posting results contains one row for each posting-model pair and the exact evidence returned by the model.",
        "Disagreements lists rows where the model differs from the reference review.",
        "Call quality shows failures or invalid outputs. Reference quality checks whether the benchmark itself is usable.",
        "The field guide defines every LLM output. UNCERTAIN is kept separate from NO.",
    ]
    for offset, note in enumerate(guidance, start=6):
        overview.cell(offset, 1, note)
        overview.cell(offset, 1).font = Font(name="Arial", size=10)
    overview["A13"] = "Model summary"
    overview["A13"].font = Font(name="Arial", size=11, bold=True, color="1F4E78")
    summary_columns = [
        "model_name", "n_scored", "seo_precision", "seo_recall", "seo_f1",
        "geo_precision", "geo_recall", "geo_f1", "group_accuracy",
        "estimated_cost_usd", "mean_latency_seconds", "possible_duty_prior_confusions",
    ]
    summary = per_model.reindex(columns=summary_columns).rename(columns=EXCEL_COLUMN_LABELS)
    for col_index, value in enumerate(summary.columns, start=1):
        overview.cell(14, col_index, value)
    for row in summary.itertuples(index=False, name=None):
        overview.append([_clean_excel_value(value) for value in row])
    for cell in overview[14]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in overview.iter_rows(min_row=15, max_row=overview.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
    for column in range(3, 10):
        for row in range(15, overview.max_row + 1):
            overview.cell(row, column).number_format = "0.0%"
    for row in range(15, overview.max_row + 1):
        overview.cell(row, 10).number_format = '"$"0.0000'
        overview.cell(row, 11).number_format = "0.00"
    overview.column_dimensions["A"].width = 36
    for column in range(2, 13):
        overview.column_dimensions[get_column_letter(column)].width = 16
    overview.row_dimensions[14].height = 42
    overview.freeze_panes = "A14"

    _write_frame_sheet(workbook, "Posting results", scored)
    _write_frame_sheet(workbook, "Disagreements", disagreements)
    _write_frame_sheet(workbook, "Call quality", call_quality)
    _write_frame_sheet(workbook, "Confusion matrices", confusion)
    _write_frame_sheet(workbook, "Reference quality", reference_quality)
    _write_frame_sheet(
        workbook,
        "Field guide",
        pd.DataFrame(FIELD_DEFINITIONS, columns=["field", "plain_language_meaning", "values_or_format"]),
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output / "pilot_results.xlsx")


def load_jsonl(path: str | Path, *, optional: bool = False) -> list[dict]:
    source = Path(path)
    if optional and not source.exists():
        return []
    records = []
    for line_number, line in enumerate(source.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON in {source} line {line_number}: {exc}") from exc
        if not isinstance(value, dict):
            raise ValueError(f"Expected a JSON object in {source} line {line_number}")
        records.append(value)
    return records


def load_table(path: str, sheet: str | None = None) -> pd.DataFrame:
    source = Path(path)
    if source.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(source, sheet_name=sheet or 0, dtype={"job_id": str})
    if source.suffix.lower() == ".parquet":
        frame = pd.read_parquet(source)
        frame["job_id"] = frame["job_id"].astype(str)
        return frame
    return pd.read_csv(source, dtype={"job_id": str})


def parse_reference(value):
    if isinstance(value, dict):
        return value
    if not isinstance(value, str):
        return {}
    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(value)
            if isinstance(parsed, dict):
                return parsed
        except (ValueError, SyntaxError, TypeError):
            pass
    return {}


def safe_div(numerator: int | float, denominator: int | float) -> float:
    return float(numerator / denominator) if denominator else 0.0


def binary_metrics(true: pd.Series, predicted: pd.Series) -> dict[str, float]:
    pairs = [(str(t), str(p)) for t, p in zip(true, predicted) if pd.notna(t) and pd.notna(p)]
    true_positive = sum(t == "YES" and p == "YES" for t, p in pairs)
    false_positive = sum(t != "YES" and p == "YES" for t, p in pairs)
    false_negative = sum(t == "YES" and p != "YES" for t, p in pairs)
    precision = safe_div(true_positive, true_positive + false_positive)
    recall = safe_div(true_positive, true_positive + false_negative)
    return {
        "precision": precision,
        "recall": recall,
        "f1": safe_div(2 * precision * recall, precision + recall),
    }


def accuracy(true: pd.Series, predicted: pd.Series) -> float:
    pairs = [(str(t), str(p)) for t, p in zip(true, predicted) if pd.notna(t) and pd.notna(p)]
    return safe_div(sum(t == p for t, p in pairs), len(pairs))


def latest_by_key(records: list[dict]) -> dict[tuple[str, str], dict]:
    latest = {}
    for record in records:
        key = (str(record.get("job_id", "")), str(record.get("model_name", "")))
        latest[key] = record
    return latest


def evidence_is_valid(quote: str, source_text: str | None) -> bool | None:
    if not quote:
        return None
    if source_text is None:
        return None
    return quote in source_text


def metric_row(group: pd.DataFrame, model_name: str, stratum: str | None = None) -> dict:
    seo = binary_metrics(group["seo_true"], group["seo_pred"])
    geo = binary_metrics(group["geo_true"], group["geo_pred"])
    row = {
        "model_name": model_name,
        "n_scored": len(group),
        "seo_precision": seo["precision"],
        "seo_recall": seo["recall"],
        "seo_f1": seo["f1"],
        "geo_precision": geo["precision"],
        "geo_recall": geo["recall"],
        "geo_f1": geo["f1"],
        "seo_accuracy": accuracy(group["seo_true"], group["seo_pred"]),
        "geo_accuracy": accuracy(group["geo_true"], group["geo_pred"]),
        "group_accuracy": accuracy(group["group_true"], group["group_pred"]),
        "estimated_cost_usd": group["estimated_cost_usd"].sum(),
        "mean_latency_seconds": group["latency_seconds"].mean(),
        "mean_attempts": group["attempts"].mean(),
        "possible_duty_prior_confusions": int(group["possible_duty_prior_confusion"].sum()),
    }
    for prefix in ("seo", "geo"):
        valid = group[f"{prefix}_evidence_valid"].dropna()
        row[f"{prefix}_evidence_substring_rate"] = valid.mean() if len(valid) else None
    if stratum is not None:
        row["pilot_stratum"] = stratum
    return row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--predictions", required=True)
    parser.add_argument("--reference", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--failures")
    parser.add_argument("--source")
    parser.add_argument("--source-sheet")
    args = parser.parse_args()

    predictions = latest_by_key(load_jsonl(args.predictions))
    failures = latest_by_key(load_jsonl(args.failures, optional=True)) if args.failures else {}
    references = {}
    origins = {}
    for record in load_jsonl(args.reference):
        job_id = str(record["job_id"])
        references[job_id] = parse_reference(record.get("reference"))
        origins[job_id] = record.get("reference_origin", "")

    source_text = {}
    source_titles = {}
    if args.source:
        source = load_table(args.source, args.source_sheet).fillna("")
        if "job_id" not in source.columns or "description" not in source.columns:
            raise ValueError("Source table must contain job_id and description")
        if source["job_id"].duplicated().any():
            raise ValueError("job_id must be unique in the source table")
        for _, record in source.iterrows():
            source_id = str(record["job_id"])
            source_text[source_id] = "\n".join(
                str(record.get(column, "")) for column in ("title_raw", "jobtitle_translated", "description")
            )
            source_titles[source_id] = {
                "title_raw": str(record.get("title_raw", "")),
                "jobtitle_translated": str(record.get("jobtitle_translated", "")),
            }

    rows = []
    for (job_id, model_name), record in predictions.items():
        reference = references.get(job_id)
        if not reference:
            continue
        prediction = record.get("prediction", {})
        combined_source = source_text.get(job_id)
        seo_valid = evidence_is_valid(str(prediction.get("seo_evidence", "")), combined_source)
        geo_valid = evidence_is_valid(str(prediction.get("geo_evidence", "")), combined_source)
        prior_reference = str(reference.get("required_prior_experience", "")).strip()
        duty_prior_proxy = (
            prior_reference != ""
            and (
                (prediction.get("seo_duty") == "YES" and reference.get("seo_duty") == "NO")
                or (prediction.get("geo_duty") == "YES" and reference.get("geo_duty") == "NO")
            )
        )
        rows.append(
            {
                "job_id": job_id,
                "title_raw": source_titles.get(job_id, {}).get("title_raw", ""),
                "jobtitle_translated": source_titles.get(job_id, {}).get("jobtitle_translated", ""),
                "reference_origin": origins.get(job_id, ""),
                "pilot_stratum": record.get("pilot_stratum", ""),
                "model_name": model_name,
                "seo_true": reference.get("seo_duty"),
                "seo_pred": prediction.get("seo_duty"),
                "geo_true": reference.get("geo_duty"),
                "geo_pred": prediction.get("geo_duty"),
                "group_true": reference.get("summary_group"),
                "group_pred": prediction.get("summary_group"),
                "seo_evidence": prediction.get("seo_evidence", ""),
                "geo_evidence": prediction.get("geo_evidence", ""),
                "seo_centrality": prediction.get("seo_centrality", ""),
                "geo_centrality": prediction.get("geo_centrality", ""),
                "required_prior_experience": prediction.get("required_prior_experience", ""),
                "prior_experience_evidence": prediction.get("prior_experience_evidence", ""),
                "seo_background_for_geo": prediction.get("seo_background_for_geo", ""),
                "adjacent_type": prediction.get("adjacent_type", ""),
                "text_completeness": prediction.get("text_completeness", ""),
                "uncertainty_reason": prediction.get("uncertainty_reason", ""),
                "concise_rationale": prediction.get("concise_rationale", ""),
                "seo_evidence_valid": seo_valid,
                "geo_evidence_valid": geo_valid,
                "possible_duty_prior_confusion": duty_prior_proxy,
                "validation_warnings": " | ".join(record.get("validation_warnings", [])),
                "evidence_substrings_valid": record.get("evidence_substrings_valid", True),
                "provider_schema_exact": record.get("provider_schema_exact", True),
                "input_tokens": record.get("input_tokens", 0),
                "output_tokens": record.get("output_tokens", 0),
                "latency_seconds": record.get("latency_seconds", 0),
                "estimated_cost_usd": record.get("estimated_cost_usd", 0),
                "attempts": record.get("attempts", 1),
            }
        )

    columns = [
        "job_id", "title_raw", "jobtitle_translated", "reference_origin", "pilot_stratum", "model_name", "seo_true", "seo_pred",
        "geo_true", "geo_pred", "group_true", "group_pred", "seo_evidence", "geo_evidence",
        "seo_centrality", "geo_centrality", "required_prior_experience", "prior_experience_evidence",
        "seo_background_for_geo", "adjacent_type", "text_completeness", "uncertainty_reason", "concise_rationale",
        "seo_evidence_valid", "geo_evidence_valid", "possible_duty_prior_confusion",
        "validation_warnings", "evidence_substrings_valid", "provider_schema_exact",
        "input_tokens", "output_tokens", "latency_seconds", "estimated_cost_usd", "attempts",
    ]
    scored = pd.DataFrame(rows, columns=columns)
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)

    reference_quality = []
    for job_id, reference in sorted(references.items()):
        schema_errors = validate_prediction(reference)
        all_errors = validate_prediction(reference, source_text.get(job_id)) if args.source else schema_errors
        evidence_errors = [error for error in all_errors if "exact substring" in error]
        reference_quality.append(
            {
                "job_id": job_id,
                "reference_origin": origins.get(job_id, ""),
                "schema_or_logic_valid": not schema_errors,
                "evidence_substrings_valid": None if not args.source else not evidence_errors,
                "schema_or_logic_errors": " | ".join(schema_errors),
                "evidence_errors": " | ".join(evidence_errors),
            }
        )
    reference_quality_frame = pd.DataFrame(reference_quality)
    reference_quality_frame.to_csv(output / "reference_quality.csv", index=False)

    per_model = []
    per_stratum = []
    if not scored.empty:
        for model_name, group in scored.groupby("model_name", sort=True):
            per_model.append(metric_row(group, model_name))
        for (model_name, stratum), group in scored.groupby(["model_name", "pilot_stratum"], sort=True):
            per_stratum.append(metric_row(group, model_name, stratum))
    pd.DataFrame(per_model).to_csv(output / "per_model_metrics.csv", index=False)
    pd.DataFrame(per_stratum).to_csv(output / "per_stratum_metrics.csv", index=False)
    pd.DataFrame(per_model).to_csv(output / "summary.csv", index=False)

    disagreements = scored[
        (scored["seo_true"] != scored["seo_pred"])
        | (scored["geo_true"] != scored["geo_pred"])
        | (scored["group_true"] != scored["group_pred"])
    ] if not scored.empty else scored
    disagreements.to_csv(output / "disagreements.csv", index=False)

    confusion_rows = []
    for model_name, group in scored.groupby("model_name", sort=True):
        for field, labels in (("seo", LABEL_ORDER), ("geo", LABEL_ORDER), ("group", GROUP_ORDER)):
            for true_label in labels:
                for predicted_label in labels:
                    count = int(
                        ((group[f"{field}_true"] == true_label) & (group[f"{field}_pred"] == predicted_label)).sum()
                    )
                    confusion_rows.append(
                        {
                            "model_name": model_name,
                            "field": field,
                            "true_label": true_label,
                            "predicted_label": predicted_label,
                            "count": count,
                        }
                    )
    confusion_frame = pd.DataFrame(confusion_rows)
    confusion_frame.to_csv(output / "confusion_matrices.csv", index=False)

    model_names = sorted({key[1] for key in predictions} | {key[1] for key in failures})
    call_summary = []
    for model_name in model_names:
        success_keys = {key for key in predictions if key[1] == model_name}
        failure_keys = {key for key in failures if key[1] == model_name and key not in success_keys}
        invalid_keys = {
            key
            for key in failure_keys
            if failures[key].get("error_type") in {"invalid_json", "schema_or_semantic_validation"}
        }
        attempted = len(success_keys) + len(failure_keys)
        call_summary.append(
            {
                "model_name": model_name,
                "successful_calls": len(success_keys),
                "terminal_failed_calls": len(failure_keys),
                "terminal_invalid_output_calls": len(invalid_keys),
                "failed_call_rate": safe_div(len(failure_keys), attempted),
                "invalid_output_rate": safe_div(len(invalid_keys), attempted),
            }
        )
    call_quality_frame = pd.DataFrame(call_summary)
    call_quality_frame.to_csv(output / "call_quality.csv", index=False)

    write_excel_workbook(
        output,
        pd.DataFrame(per_model),
        scored,
        disagreements,
        call_quality_frame,
        confusion_frame,
        reference_quality_frame,
    )

    summary = {
        "rows_scored": len(scored),
        "reference_jobs": len(references),
        "models": model_names,
        "source_text_supplied_for_evidence_check": bool(args.source),
        "prediction_records_after_deduplication": len(predictions),
        "terminal_failure_records_after_deduplication": len(failures),
        "reference_schema_or_logic_invalid": int(
            (~reference_quality_frame["schema_or_logic_valid"]).sum()
        ) if not reference_quality_frame.empty else 0,
        "reference_evidence_substring_invalid": int(
            (reference_quality_frame["evidence_substrings_valid"] == False).sum()
        ) if args.source and not reference_quality_frame.empty else None,
        "possible_duty_prior_confusion_is_proxy": True,
        "notes": [
            "Evidence substring validity checks literal occurrence only; it does not prove semantic support.",
            "Possible duty/prior-experience confusion is a review flag, not a definitive semantic error.",
            "Precision and recall treat YES as the positive class and NO/UNCERTAIN as non-positive.",
        ],
        "call_quality": call_summary,
    }
    (output / "evaluation.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Wrote evaluation to {output}")


if __name__ == "__main__":
    main()
