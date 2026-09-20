from __future__ import annotations

import json
import os
import tempfile
import unittest
import sys
from argparse import Namespace
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from revelio_pilot import providers
from revelio_pilot import evaluate
from revelio_pilot import run_pilot
from revelio_pilot.prepare import _boolean_series, existing
from revelio_pilot.prompt import get_system_prompt
from revelio_pilot.schema import REQUIRED_FIELDS, normalize_prediction, validate_prediction


def valid_prediction(**updates):
    value = {
        "seo_duty": "YES", "geo_duty": "NO", "seo_centrality": "PRIMARY",
        "geo_centrality": "NOT_APPLICABLE", "seo_evidence": "Improve organic search rankings",
        "geo_evidence": "", "required_prior_experience": "three years",
        "prior_experience_evidence": "three years of SEO experience",
        "seo_background_for_geo": "NO_EVIDENCE", "adjacent_type": "",
        "text_completeness": "FULL", "uncertainty_reason": "",
        "summary_group": "SEO_ONLY", "concise_rationale": "SEO is an assigned duty; GEO is not.",
    }
    value.update(updates)
    return value


class SchemaTests(unittest.TestCase):
    def test_valid_prediction_and_evidence(self):
        source = "Improve organic search rankings; requires three years of SEO experience."
        self.assertEqual(validate_prediction(valid_prediction(), source), [])

    def test_rejects_extra_fields_and_inconsistent_group(self):
        errors = " | ".join(validate_prediction(valid_prediction(summary_group="BOTH", invented=True)))
        self.assertIn("unexpected fields", errors)
        self.assertIn("summary_group must be SEO_ONLY", errors)

    def test_requires_exact_source_quote(self):
        errors = validate_prediction(valid_prediction(), "different source")
        self.assertTrue(any("exact substring" in error for error in errors))

    def test_rejects_evidence_when_duty_is_no(self):
        errors = validate_prediction(valid_prediction(geo_evidence="AI search"))
        self.assertIn("geo_evidence must be empty when duty is NO", errors)

    def test_schema_field_set_is_exact(self):
        self.assertEqual(set(valid_prediction()), set(REQUIRED_FIELDS))

    def test_normalization_removes_and_reports_provider_added_fields(self):
        normalized, warnings = normalize_prediction(
            valid_prediction(geo_background_for_geo="NO_EVIDENCE")
        )
        self.assertEqual(set(normalized), set(REQUIRED_FIELDS))
        self.assertEqual(
            warnings,
            ["removed unexpected fields: ['geo_background_for_geo']"],
        )


class PrepareTests(unittest.TestCase):
    def test_string_false_is_not_truthy(self):
        result = _boolean_series(pd.Series(["false", "TRUE", "0", "yes"]))
        self.assertEqual(result.tolist(), [False, True, False, True])

    def test_existing_samples_only_reviewed_ids(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            candidates = pd.DataFrame({
                "job_id": [str(i) for i in range(6)], "title_raw": ["t"] * 6,
                "jobtitle_translated": ["t"] * 6, "description": ["d"] * 6,
            })
            references = pd.DataFrame([
                {"job_id": "1", **valid_prediction()}, {"job_id": "3", **valid_prediction()},
            ])
            candidates.to_csv(root / "candidates.csv", index=False)
            references.to_csv(root / "references.csv", index=False)
            args = Namespace(
                input=str(root / "candidates.csv"), sheet=None, n=2, seed=7,
                reference_input=str(root / "references.csv"), reference_sheet=None,
                reference_col="assistant_final_judgment", reference_origin="test",
                reference_output=str(root / "private" / "labels.jsonl"), output=str(root / "sample.csv"),
            )
            existing(args)
            sampled = pd.read_csv(root / "sample.csv", dtype={"job_id": str})
            self.assertEqual(set(sampled["job_id"]), {"1", "3"})
            self.assertEqual(set(sampled["pilot_stratum"]), {"REVIEWED_EXISTING"})
            records = [json.loads(line) for line in (root / "private" / "labels.jsonl").read_text().splitlines()]
            self.assertEqual({record["job_id"] for record in records}, {"1", "3"})


class ProviderPayloadTests(unittest.TestCase):
    def test_openai_uses_strict_schema_and_no_storage(self):
        body = {"output": [{"content": [{"type": "output_text", "text": "{}"}]}], "usage": {}}
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test"}), patch.object(
            providers, "_post", return_value=body
        ) as post:
            providers.call_openai("model", "system", "user")
        payload = post.call_args.kwargs["payload"]
        self.assertFalse(payload["store"])
        self.assertTrue(payload["text"]["format"]["strict"])

    def test_openai_passes_explicit_reasoning_effort(self):
        body = {"output": [{"content": [{"type": "output_text", "text": "{}"}]}], "usage": {}}
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test"}), patch.object(
            providers, "_post", return_value=body
        ) as post:
            providers.call_openai(
                "gpt-5.6-luna", "system", "user", reasoning_effort="max"
            )
        self.assertEqual(
            post.call_args.kwargs["payload"]["reasoning"], {"effort": "max"}
        )

    def test_openai_preserves_usage_when_reasoning_consumes_all_output(self):
        body = {
            "output": [],
            "usage": {"input_tokens": 100, "output_tokens": 6000},
        }
        with patch.dict(os.environ, {"OPENAI_API_KEY": "test"}), patch.object(
            providers, "_post", return_value=body
        ):
            text, usage, raw = providers.call_openai(
                "gpt-5.6-luna", "system", "user", max_output_tokens=6000,
                reasoning_effort="max",
            )
        self.assertEqual(text, "")
        self.assertEqual(usage["output_tokens"], 6000)
        self.assertEqual(raw["output"], [])

    def test_provider_output_limits_use_native_fields(self):
        openai_body = {"output": [{"content": [{"type": "output_text", "text": "{}"}]}], "usage": {}}
        mistral_body = {"choices": [{"message": {"content": "{}"}}], "usage": {}}
        gemini_body = {"candidates": [{"content": {"parts": [{"text": "{}"}]}}], "usageMetadata": {}}
        with patch.dict(os.environ, {
            "OPENAI_API_KEY": "test", "MISTRAL_API_KEY": "test", "GEMINI_API_KEY": "test"
        }), patch.object(providers, "_post", side_effect=[openai_body, mistral_body, gemini_body]) as post:
            providers.call_openai("model", "system", "user", max_output_tokens=1000)
            providers.call_mistral("model", "system", "user", max_output_tokens=1000)
            providers.call_gemini("model", "system", "user", max_output_tokens=1000)
        self.assertEqual(post.call_args_list[0].kwargs["payload"]["max_output_tokens"], 1000)
        self.assertEqual(post.call_args_list[1].kwargs["payload"]["max_tokens"], 1000)
        self.assertEqual(
            post.call_args_list[2].kwargs["payload"]["generationConfig"]["maxOutputTokens"], 1000
        )

    def test_gemini_uses_generate_content_structured_output_fields(self):
        body = {"candidates": [{"content": {"parts": [{"text": "{}"}]}}], "usageMetadata": {}}
        with patch.dict(os.environ, {"GEMINI_API_KEY": "test"}), patch.object(
            providers, "_post", return_value=body
        ) as post:
            providers.call_gemini("model", "system", "user")
        generation_config = post.call_args.kwargs["payload"]["generationConfig"]
        self.assertEqual(generation_config["responseMimeType"], "application/json")
        self.assertEqual(generation_config["responseJsonSchema"], providers.PREDICTION_SCHEMA)
        self.assertNotIn("responseFormat", generation_config)

    def test_rejects_markdown_escaped_or_whitespace_keys_before_http(self):
        for bad_key in ("sk-test\\_escaped", "key with spaces", "key\nwith-newline"):
            with self.subTest(bad_key=repr(bad_key)), patch.dict(
                os.environ, {"OPENAI_API_KEY": bad_key}
            ), patch.object(providers, "_post") as post:
                with self.assertRaises(RuntimeError):
                    providers.call_openai("model", "system", "user")
                post.assert_not_called()

    def test_missing_key_fails_before_http(self):
        with patch.dict(os.environ, {}, clear=True), patch.object(providers, "_post") as post:
            with self.assertRaises(RuntimeError):
                providers.call_mistral("model", "system", "user")
        post.assert_not_called()


class EvaluationTests(unittest.TestCase):
    def test_evaluator_writes_metrics_and_confusion_matrix(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            prediction = {
                "job_id": "1", "pilot_stratum": "new_random", "model_name": "m",
                "prediction": valid_prediction(), "estimated_cost_usd": 0.01,
                "latency_seconds": 0.2, "attempts": 1,
            }
            reference = {"job_id": "1", "reference_origin": "human_adjudicated", "reference": valid_prediction()}
            (root / "predictions.jsonl").write_text(json.dumps(prediction) + "\n", encoding="utf-8")
            (root / "reference.jsonl").write_text(json.dumps(reference) + "\n", encoding="utf-8")
            pd.DataFrame([{
                "job_id": "1", "title_raw": "", "jobtitle_translated": "",
                "description": "Improve organic search rankings; three years of SEO experience",
            }]).to_csv(root / "source.csv", index=False)
            argv = [
                "evaluate", "--predictions", str(root / "predictions.jsonl"),
                "--reference", str(root / "reference.jsonl"), "--source", str(root / "source.csv"),
                "--output", str(root / "evaluation"),
            ]
            with patch.object(sys, "argv", argv):
                evaluate.main()
            metrics = pd.read_csv(root / "evaluation" / "per_model_metrics.csv")
            self.assertEqual(metrics.loc[0, "group_accuracy"], 1.0)
            self.assertTrue((root / "evaluation" / "confusion_matrices.csv").exists())
            quality = pd.read_csv(root / "evaluation" / "reference_quality.csv")
            self.assertTrue(quality.loc[0, "schema_or_logic_valid"])
            workbook = load_workbook(root / "evaluation" / "pilot_results.xlsx", data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Overview", "Posting results", "Disagreements", "Call quality", "Confusion matrices", "Reference quality", "Field guide"],
            )
            self.assertEqual(workbook["Posting results"]["A2"].value, "1")
            self.assertEqual(workbook["Field guide"]["A2"].value, "seo_duty")


class RunnerIntegrationTests(unittest.TestCase):
    @staticmethod
    def _write_runner_fixture(root: Path, *, max_retries: int = 0) -> list[str]:
        pd.DataFrame([{
            "job_id": "1", "title_raw": "SEO role", "jobtitle_translated": "SEO role",
            "description": "Improve organic search rankings; three years of SEO experience",
            "pilot_stratum": "existing_instruction_check",
        }]).to_csv(root / "input.csv", index=False)
        (root / "config.yaml").write_text(
            f"seed: 7\nmax_retries: {max_retries}\ntimeout_seconds: 10\nmax_output_tokens: 1000\n"
            "max_cost_usd: 1\nmodels:\n"
            "  - name: mock_openai\n    provider: openai\n    model_id: test-model\n"
            "    input_usd_per_million: 0.2\n    output_usd_per_million: 1.2\n",
            encoding="utf-8",
        )
        return [
            "run", "--input", "input.csv", "--config", "config.yaml",
            "--phase", "instruction_check", "--run-id", "test-run",
        ]

    @staticmethod
    def _run_in(root: Path, argv: list[str], call_mock) -> None:
        prior = Path.cwd()
        try:
            os.chdir(root)
            with patch.object(sys, "argv", argv), patch.dict(
                os.environ, {"OPENAI_API_KEY": "test"}
            ), patch.object(run_pilot, "call", call_mock):
                run_pilot.main()
        finally:
            os.chdir(prior)

    def test_gemini_cost_tokens_include_thinking_tokens(self):
        self.assertEqual(
            run_pilot.usage_tokens(
                "gemini",
                {
                    "promptTokenCount": 120,
                    "candidatesTokenCount": 40,
                    "thoughtsTokenCount": 15,
                },
            ),
            (120, 55),
        )

    def test_run_lock_rejects_a_second_process_handle(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / ".run.lock"
            first = run_pilot.RunLock(path)
            second = run_pilot.RunLock(path)
            first.acquire()
            try:
                with self.assertRaises(RuntimeError):
                    second.acquire()
            finally:
                first.release()
            second.acquire()
            second.release()

    def test_jsonl_tail_repair_is_limited_to_final_record(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "records.jsonl"
            path.write_bytes(b'{"job_id":"1"}\n{"job_id":')
            self.assertEqual(
                run_pilot.repair_jsonl_tail(path),
                "removed_incomplete_final_record",
            )
            self.assertEqual(run_pilot.load_jsonl(path), [{"job_id": "1"}])

            path.write_bytes(b'{"job_id":"1"}')
            self.assertEqual(
                run_pilot.repair_jsonl_tail(path),
                "added_missing_final_newline",
            )
            self.assertTrue(path.read_bytes().endswith(b"\n"))

            path.write_bytes(b'{"job_id":"1"}\nnot-json\n{"job_id":"2"}\n')
            self.assertIsNone(run_pilot.repair_jsonl_tail(path))
            with self.assertRaises(ValueError):
                run_pilot.load_jsonl(path)

    def test_mock_provider_run_checkpoints_validated_result(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root, max_retries=1)
            self._run_in(
                root,
                argv,
                Mock(side_effect=[
                    ("not-json", {"input_tokens": 100, "output_tokens": 50}, {"id": "mock-1"}, 0.01),
                    (json.dumps(valid_prediction()), {"input_tokens": 100, "output_tokens": 50}, {"id": "mock-2"}, 0.01),
                ]),
            )
            manifest = json.loads((root / "outputs" / "runs" / "test-run" / "manifest.json").read_text())
            self.assertEqual(manifest["successful_calls"], 1)
            self.assertFalse(manifest["stopped_for_budget"])
            raw_lines = (root / "outputs" / "runs" / "test-run" / "raw_responses.jsonl").read_text().splitlines()
            self.assertEqual(len(raw_lines), 2)
            self.assertAlmostEqual(manifest["estimated_cost_usd"], 0.00016)
            result = json.loads((root / "outputs" / "runs" / "test-run" / "predictions.jsonl").read_text())
            self.assertEqual(result["attempts"], 2)

    def test_runner_freezes_selected_prompt_version(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root) + ["--prompt-version", "improved_v2"]
            call_mock = Mock(return_value=(
                json.dumps(valid_prediction()),
                {"input_tokens": 100, "output_tokens": 50},
                {"id": "mock-v2"},
                0.01,
            ))
            self._run_in(root, argv, call_mock)
            run_dir = root / "outputs" / "runs" / "test-run"
            manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["prompt_version"], "improved_v2")
            self.assertEqual(
                (run_dir / "prompts" / "prompt.txt").read_text(encoding="utf-8"),
                get_system_prompt("improved_v2"),
            )
            self.assertEqual(call_mock.call_args.args[2], get_system_prompt("improved_v2"))

    def test_resume_preserves_retry_count_and_cost(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root, max_retries=1)
            first_call = Mock(return_value=(
                "not-json", {"input_tokens": 100, "output_tokens": 50}, {"id": "mock-1"}, 0.01,
            ))
            prior = Path.cwd()
            try:
                os.chdir(root)
                with patch.object(sys, "argv", argv), patch.dict(
                    os.environ, {"OPENAI_API_KEY": "test"}
                ), patch.object(run_pilot, "call", first_call), patch.object(
                    run_pilot.time, "sleep", side_effect=KeyboardInterrupt
                ), self.assertRaises(KeyboardInterrupt):
                    run_pilot.main()
            finally:
                os.chdir(prior)

            second_call = Mock(return_value=(
                json.dumps(valid_prediction()),
                {"input_tokens": 100, "output_tokens": 50},
                {"id": "mock-2"},
                0.01,
            ))
            self._run_in(root, argv, second_call)
            first_call.assert_called_once()
            second_call.assert_called_once()
            run_dir = root / "outputs" / "runs" / "test-run"
            result = json.loads((run_dir / "predictions.jsonl").read_text())
            self.assertEqual(result["attempts"], 2)
            self.assertAlmostEqual(result["estimated_cost_usd"], 0.00016)
            self.assertEqual(len((run_dir / "raw_responses.jsonl").read_text().splitlines()), 2)

    def test_resume_recovers_terminal_checkpoint_without_new_api_call(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root)
            real_append = run_pilot.append_jsonl

            def interrupt_prediction_append(path: Path, record: dict) -> None:
                if path.name == "predictions.jsonl":
                    raise KeyboardInterrupt
                real_append(path, record)

            prior = Path.cwd()
            try:
                os.chdir(root)
                with patch.object(sys, "argv", argv), patch.dict(
                    os.environ, {"OPENAI_API_KEY": "test"}
                ), patch.object(run_pilot, "call", return_value=(
                    json.dumps(valid_prediction()),
                    {"input_tokens": 100, "output_tokens": 50},
                    {"id": "mock-1"},
                    0.01,
                )), patch.object(
                    run_pilot, "append_jsonl", side_effect=interrupt_prediction_append
                ), self.assertRaises(KeyboardInterrupt):
                    run_pilot.main()
            finally:
                os.chdir(prior)

            second_call = Mock()
            self._run_in(root, argv, second_call)
            second_call.assert_not_called()
            run_dir = root / "outputs" / "runs" / "test-run"
            self.assertEqual(len((run_dir / "predictions.jsonl").read_text().splitlines()), 1)
            manifest = json.loads((run_dir / "manifest.json").read_text())
            self.assertEqual(manifest["run_status"], "completed")

    def test_resume_does_not_duplicate_ambiguous_in_flight_call(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root)
            with self.assertRaises(KeyboardInterrupt):
                self._run_in(root, argv, Mock(side_effect=KeyboardInterrupt))

            second_call = Mock()
            with self.assertRaises(RuntimeError):
                self._run_in(root, argv, second_call)
            second_call.assert_not_called()
            run_dir = root / "outputs" / "runs" / "test-run"
            self.assertFalse((run_dir / "failures.jsonl").exists())

            accept_call = Mock()
            self._run_in(
                root,
                argv + ["--accept-ambiguous-in-flight-as-failure"],
                accept_call,
            )
            accept_call.assert_not_called()
            failure = json.loads((run_dir / "failures.jsonl").read_text())
            self.assertEqual(failure["error_type"], "interrupted_in_flight")
            self.assertTrue(failure["cost_unknown"])

    def test_ambiguous_in_flight_retry_requires_explicit_flag(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            argv = self._write_runner_fixture(root)
            with self.assertRaises(KeyboardInterrupt):
                self._run_in(root, argv, Mock(side_effect=KeyboardInterrupt))

            retry_call = Mock(return_value=(
                json.dumps(valid_prediction()),
                {"input_tokens": 100, "output_tokens": 50},
                {"id": "replacement"},
                0.01,
            ))
            self._run_in(root, argv + ["--retry-ambiguous-in-flight"], retry_call)
            retry_call.assert_called_once()
            run_dir = root / "outputs" / "runs" / "test-run"
            result = json.loads((run_dir / "predictions.jsonl").read_text())
            self.assertEqual(result["ambiguous_in_flight_attempts"], 1)
            self.assertTrue(result["cost_unknown"])


if __name__ == "__main__":
    unittest.main()
